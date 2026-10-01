"""
Audit what has actually been saved: cross-check the workbook, the JSON and the JSONL
against each other and report any disagreement.

    python verify_data.py           # full audit
    python verify_data.py --show 20 # more examples per finding

The JSONL is the reference, because it is appended the instant an answer arrives and is
never rewritten. Every record in it should correspond to exactly one workbook cell, in
the row whose Query column holds the question that was actually asked. That last point is
what catches the failure that would matter most: an answer landing on the wrong row.

Reading is safe while a run is in progress; an answer captured seconds ago may legitimately
not be in the workbook yet, and that is reported separately from a real mismatch.
"""

import argparse
import json
import pathlib
import sys

import openpyxl

import llm_engines
import run_tracker

HERE = pathlib.Path(__file__).parent
DEFAULT_WORKBOOK = "LLM_Response_Trackers.xlsx"


def norm(text, engine=None):
    """Compare fairly: strip the UI noise and collapse whitespace."""
    out = str(text or "")
    if engine is not None:
        out = llm_engines.clean(engine, out)
    return " ".join(out.split())


def main():
    ap = argparse.ArgumentParser(description="Verify the saved answers.")
    ap.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    ap.add_argument("--show", type=int, default=6)
    args = ap.parse_args()

    path = pathlib.Path(args.workbook)
    if not path.is_absolute():
        path = HERE / path
    json_path = path.with_suffix(".json")
    jsonl_path = path.with_name(path.stem + "_answers.jsonl")

    if not path.exists():
        sys.exit(f"Workbook not found: {path}")
    print(f"workbook : {path.name}  ({path.stat().st_size:,} bytes)")
    print(f"json     : {json_path.name}  "
          f"({json_path.stat().st_size:,} bytes)" if json_path.exists() else "json  : MISSING")
    print(f"jsonl    : {jsonl_path.name}  "
          f"({jsonl_path.stat().st_size:,} bytes)" if jsonl_path.exists() else "jsonl : MISSING")

    records = []
    if jsonl_path.exists():
        for n, line in enumerate(jsonl_path.read_text(encoding="utf-8").splitlines(), 1):
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except Exception as exc:
                print(f"  BAD JSONL LINE {n}: {exc}")
    print(f"\n{len(records)} JSONL records to check")

    wb = openpyxl.load_workbook(path)
    by_label = {e.label: e for e in llm_engines.ENGINES.values()}

    # Map every answer cell in the workbook: (sheet, row, engine) -> (value, query, cat, src)
    cells = {}
    for ws in wb.worksheets:
        if ws.title == "Index":
            continue
        for section in run_tracker.find_sections(ws):
            for row in section["rows"]:
                query = str(ws.cell(row, section["query_col"]).value or "")
                category = str(ws.cell(row, 2).value or "")
                for label, col in section["out_cols"].items():
                    src_col = section["src_cols"].get(label)
                    cells[(ws.title, row, label)] = (
                        ws.cell(row, col).value,
                        query, category, section["tier"],
                        ws.cell(row, src_col).value if src_col else None,
                    )

    findings = {
        "missing": [], "mismatch": [], "wrong_query": [], "wrong_category": [],
        "no_cell": [], "empty": [], "sources_lost": [], "truncated": [],
        "orphan_cells": [], "dupes": [],
    }

    seen = {}
    for rec in records:
        key = (rec["sub_industry"], rec["row"], rec["engine"])
        seen.setdefault(key, []).append(rec)
        if key not in cells:
            findings["no_cell"].append(f"{key} -- no such row/engine in the workbook")
            continue
        value, query, category, tier, src = cells[key]
        engine = by_label.get(rec["engine"])

        if value is None or not str(value).strip():
            findings["missing"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']} {rec['engine']}: "
                f"{rec['chars']:,}ch in the log, cell is empty")
            continue

        cell_text = norm(value, engine)
        log_text = norm(rec["answer"], engine)
        trimmed = "trimmed for Excel" in str(value)
        if trimmed:
            findings["truncated"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']} {rec['engine']}: "
                f"{rec['chars']:,}ch -> cell trimmed (full text in JSONL/JSON)")
        elif cell_text[:400] != log_text[:400]:
            findings["mismatch"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']} {rec['engine']}:\n"
                f"       cell: {cell_text[:90]!r}\n"
                f"        log: {log_text[:90]!r}")

        # The alignment check that matters: is this answer on the row that asked it?
        if norm(rec["query"]) != norm(query):
            findings["wrong_query"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']} {rec['engine']}:\n"
                f"       row asks : {norm(query)[:88]!r}\n"
                f"       log asked: {norm(rec['query'])[:88]!r}")
        if norm(rec["category"]) != norm(category):
            findings["wrong_category"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']}: row={norm(category)[:40]!r} "
                f"log={norm(rec['category'])[:40]!r}")
        if rec.get("source_count") and not (src and str(src).strip()):
            findings["sources_lost"].append(
                f"{rec['sub_industry'][:20]} r{rec['row']} {rec['engine']}: "
                f"{rec['source_count']} sources logged, Sources cell empty")

    for key, group in seen.items():
        if len(group) > 1:
            findings["dupes"].append(f"{key} appears {len(group)}x in the JSONL")

    logged = set(seen)
    for key, (value, *_rest) in cells.items():
        if value is not None and str(value).strip() and key not in logged:
            findings["orphan_cells"].append(key)

    labels = {
        "missing": "answers in the log but NOT in the workbook (may be a pending save)",
        "mismatch": "cell text disagrees with the logged answer",
        "wrong_query": "answer sitting on a row that asked a DIFFERENT question",
        "wrong_category": "category disagrees between row and log",
        "no_cell": "log refers to a row/engine that does not exist",
        "empty": "empty answers saved",
        "sources_lost": "sources logged but the Sources cell is empty",
        "truncated": "answers trimmed to fit Excel (full text kept in JSON/JSONL)",
        "orphan_cells": "filled cells with no JSONL record (expected: your manual rows)",
        "dupes": "the same (row, engine) logged more than once",
    }
    print("\n" + "=" * 74)
    bad = 0
    for key, label in labels.items():
        items = findings[key]
        benign = key in ("orphan_cells", "truncated", "missing")
        flag = "OK  " if not items else ("note" if benign else "FAIL")
        if items and not benign:
            bad += len(items)
        print(f"  {flag}  {len(items):>5}  {label}")
    print("=" * 74)

    for key, label in labels.items():
        items = findings[key]
        if not items:
            continue
        if key == "orphan_cells":
            print(f"\n{label}: {len(items)}")
            sheets = {}
            for sheet, row, eng in items:
                sheets[sheet] = sheets.get(sheet, 0) + 1
            for sheet, n in sheets.items():
                print(f"    {sheet}: {n} cells")
            continue
        print(f"\n{label}:")
        for item in items[: args.show]:
            print(f"    {item}")
        if len(items) > args.show:
            print(f"    ... and {len(items) - args.show} more")

    if json_path.exists():
        data = json.loads(json_path.read_text(encoding="utf-8"))
        wb_counts = {}
        for (sheet, row, label), (value, *_r) in cells.items():
            if value is not None and str(value).strip():
                wb_counts[label] = wb_counts.get(label, 0) + 1
        print(f"\nJSON totals   : {data['totals']['answers']}")
        print(f"workbook count: {wb_counts}")
        print(f"JSON queries  : {data['totals']['queries']} "
              f"(workbook rows: {len(cells) // max(1, len(by_label))})")
        if data["totals"]["answers"] != wb_counts:
            print("  note: the JSON is rewritten on each save, so it can lag the "
                  "workbook by one answer during a live run")

    print("\nVERDICT:", "no integrity problems found" if not bad
          else f"{bad} problem(s) above need attention")


if __name__ == "__main__":
    main()
