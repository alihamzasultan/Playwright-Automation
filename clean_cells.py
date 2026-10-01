"""
Strip interface text out of answers that were captured before the cleanup existed.

Grok prefixes its answers with its own timer ("Worked for 13s ...") and Claude's message
element carries a screen-reader label ("Claude responded: ..."). Both are now stripped at
capture time, but rows saved earlier still have them. This applies the same rules to what
is already in the workbook.

    python clean_cells.py                 # report what would change, write nothing
    python clean_cells.py --apply         # rewrite the cells and refresh the JSON

Idempotent: the patterns only match at the very start of a cell, so running it twice
changes nothing the second time, and an answer that merely mentions the phrase
mid-sentence is left alone.
"""

import argparse
import pathlib
import sys

import openpyxl

import llm_engines
import run_tracker

HERE = pathlib.Path(__file__).parent
DEFAULT_WORKBOOK = "LLM_Response_Trackers.xlsx"


def main():
    ap = argparse.ArgumentParser(description="Strip UI noise from saved answers.")
    ap.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    ap.add_argument("--apply", action="store_true", help="write the changes")
    ap.add_argument("--show", type=int, default=12, help="examples to print")
    args = ap.parse_args()

    path = pathlib.Path(args.workbook)
    if not path.is_absolute():
        path = HERE / path
    if not path.exists():
        sys.exit(f"Workbook not found: {path}")
    if args.apply:
        run_tracker.check_writable(path)

    by_label = {e.label: e for e in llm_engines.ENGINES.values()}
    wb = openpyxl.load_workbook(path)

    changed, examples, scanned = 0, [], 0
    for ws in wb.worksheets:
        if ws.title == "Index":
            continue
        for section in run_tracker.find_sections(ws):
            for label, col in section["out_cols"].items():
                engine = by_label.get(label)
                if not engine or not engine.noise:
                    continue  # ChatGPT and AI Overviews have no noise patterns
                for row in section["rows"]:
                    value = ws.cell(row, col).value
                    if value is None or not str(value).strip():
                        continue
                    scanned += 1
                    before = str(value)
                    after = llm_engines.clean(engine, before)
                    if after == before:
                        continue
                    changed += 1
                    if len(examples) < args.show:
                        removed = len(before) - len(after)
                        examples.append(
                            f"  {ws.title[:22]:<23} r{row:<4} {label:<8} "
                            f"-{removed:>3}ch  {before[:34]!r} -> {after[:30]!r}")
                    if args.apply:
                        ws.cell(row, col, after)

    print(f"\nscanned {scanned} answer cells for the engines that have noise patterns")
    print(f"{changed} cell(s) {'cleaned' if args.apply else 'would be cleaned'}\n")
    for line in examples:
        print(line)
    if changed > len(examples):
        print(f"  ... and {changed - len(examples)} more")

    if not changed:
        print("Nothing to do.")
        return
    if not args.apply:
        print("\nDry run. Re-run with --apply to write these changes.")
        return

    run_tracker.save_atomically(path, lambda p: wb.save(p))
    json_path = path.with_suffix(".json")
    run_tracker.write_json(wb, path, json_path)
    print(f"\nSaved {path.name} and refreshed {json_path.name}")
    print("The JSONL log is left untouched on purpose -- it is the raw capture record.")


if __name__ == "__main__":
    main()
