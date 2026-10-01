"""
Work through the tracker workbook automatically across every LLM.

For each query with an empty answer cell, the query is put to each selected engine
(ChatGPT, Claude, Google AI Overviews, Grok) in a real browser and the answer is written
straight back into that engine's '<Engine> - Output' column. A JSON file is refreshed
after every answer, and every answer is also appended to a JSONL log the moment it
arrives.

Resumable by design: only cells that are still empty are attempted, so stopping it
(Ctrl+C) and running it again continues where it left off and never repeats a query.

    python run_tracker.py --dry-run                          # what would run
    python run_tracker.py --engines chatgpt --sheet Hospitals --limit 3
    python run_tracker.py --engines all --sheet Hospitals     # one sub-industry, 4 LLMs
    python run_tracker.py                                    # everything still pending

Each engine needs to be signed in once in the profile you run with:

    python chatgpt_query.py --login        # opens all four sites in a plain Chrome
    python llm_engines.py --probe all      # checks each one answers before a long run
"""

import argparse
import datetime
import json
import os
import pathlib
import random
import sys
import time

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import chatgpt_query as cgq
import llm_engines

HERE = pathlib.Path(__file__).parent
DEFAULT_WORKBOOK = "LLM_Response_Trackers.xlsx"
TIER_SECTION_SUFFIX = "Buying-Intent AI Queries"
QUERY_HEADER = "Query"
OUTPUT_SUFFIX = " - Output"
SOURCE_SUFFIX = " - Sources"

LIMIT_PHRASES = (
    "you've reached your limit",
    "you have reached our limit",
    "usage limit",
    "rate limit",
    "too many requests",
    "please wait a moment",
    "unusual activity",
    "upgrade to continue",
)


def log(msg):
    print(f"[{datetime.datetime.now():%H:%M:%S}] {msg}", flush=True)


# ----------------------------------------------------------------------------------
# Reading the workbook
# ----------------------------------------------------------------------------------
def find_sections(ws):
    """
    -> [{tier, header_row, query_col, out_cols: {engine label: col}, rows: [...]}]

    A tier section is a row whose column B ends with 'Buying-Intent AI Queries'; the next
    row holds the column headers, and the data runs until the next blank query cell.
    """
    sections = []
    row = 1
    while row <= ws.max_row:
        marker = str(ws.cell(row, 2).value or "").strip()
        if not marker.endswith(TIER_SECTION_SUFFIX):
            row += 1
            continue

        header_row = row + 1
        query_col, out_cols = None, {}
        for col in range(1, ws.max_column + 1):
            name = str(ws.cell(header_row, col).value or "").strip()
            if name == QUERY_HEADER:
                query_col = col
            elif name.endswith(OUTPUT_SUFFIX):
                out_cols[name[: -len(OUTPUT_SUFFIX)]] = col
        src_cols = {}
        for col in range(1, ws.max_column + 1):
            name = str(ws.cell(header_row, col).value or "").strip()
            if name.endswith(SOURCE_SUFFIX):
                src_cols[name[: -len(SOURCE_SUFFIX)]] = col
        if not query_col or not out_cols:
            row += 1
            continue

        rows = []
        r = header_row + 1
        while r <= ws.max_row:
            value = ws.cell(r, query_col).value
            if value is None or not str(value).strip():
                break
            rows.append(r)
            r += 1

        sections.append({"tier": marker, "header_row": header_row,
                         "query_col": query_col, "out_cols": out_cols,
                         "src_cols": src_cols, "rows": rows})
        row = r
    return sections


def pick_sheets(wb, sheet_filter):
    """
    Exact match wins, then prefix, then substring: plain "Hospitals" should mean
    'Hospitals & Health Systems', not also 'Veterinary Clinics & Hospitals'.
    """
    titles = [ws.title for ws in wb.worksheets if ws.title != "Index"]
    if not sheet_filter:
        return titles
    want = sheet_filter.strip().lower()
    for test in (lambda t: t.lower() == want,
                 lambda t: t.lower().startswith(want),
                 lambda t: want in t.lower()):
        hits = [t for t in titles if test(t)]
        if hits:
            return hits
    return []


def collect_tasks(wb, engines, sheet_filter=None, tier_filter=None):
    """
    -> (tasks, totals). One task per empty (query, engine) cell, ordered row by row so a
    query is finished across every engine before moving on.
    """
    wanted = set(pick_sheets(wb, sheet_filter))
    labels = [e.label for e in engines]
    tasks = []
    totals = {"rows": 0, "cells": 0, "filled": 0, "missing_cols": set()}

    for ws in wb.worksheets:
        if ws.title == "Index" or ws.title not in wanted:
            continue
        for section in find_sections(ws):
            if tier_filter and tier_filter.lower() not in section["tier"].lower():
                continue
            for label in labels:
                if label not in section["out_cols"]:
                    totals["missing_cols"].add(f"{ws.title}: {label}")
            for row in section["rows"]:
                totals["rows"] += 1
                query = str(ws.cell(row, section["query_col"]).value).strip()
                category = str(ws.cell(row, 2).value or "").strip()
                for label in labels:
                    col = section["out_cols"].get(label)
                    if not col:
                        continue
                    totals["cells"] += 1
                    existing = ws.cell(row, col).value
                    if existing is not None and str(existing).strip():
                        totals["filled"] += 1
                        continue
                    tasks.append({
                        "sheet": ws.title, "tier": section["tier"], "row": row,
                        "category": category, "query": query,
                        "engine": label, "out_col": col,
                        "src_col": section["src_cols"].get(label),
                    })
    return tasks, totals


# ----------------------------------------------------------------------------------
# Saving: workbook, nested JSON, append-only log
# ----------------------------------------------------------------------------------
def save_atomically(path, write):
    """
    Write through a temp file and swap it in, so an interruption cannot leave a
    half-written file where the only copy of the answers used to be.
    """
    tmp = path.with_suffix(path.suffix + ".tmp")
    write(tmp)
    os.replace(tmp, path)


INDEX_FIXED_HEADERS = ["#", "Category", "Sub-Industry", "Sheet", "Categories",
                       "Queries", "Rows Answered"]


def index_headers():
    """
    The per-engine columns must be named in the same order the counts are written in,
    which is the order of llm_engines.ENGINES -- spelling them out by hand once put
    Grok's count under the 'AI Overviews' heading and vice versa.
    """
    return (INDEX_FIXED_HEADERS
            + [e.label for e in llm_engines.ENGINES.values()] + ["Done %"])


def refresh_index(wb):
    """
    Recompute the Index sheet from the answer columns.

    build_trackers.py writes the Index once, so its counts freeze at whatever was true
    when the workbook was created -- which is why it kept saying 26 for Hospitals long
    after the runner had filled 37 rows. This recounts every sheet and rewrites the
    figures, including a per-engine breakdown and a total row.
    """
    if "Index" not in wb.sheetnames:
        return 0
    index = wb["Index"]
    labels = [e.label for e in llm_engines.ENGINES.values()]

    # Header row 3, as written by build_trackers.py.
    for col, name in enumerate(index_headers(), start=1):
        cell = index.cell(3, col, name)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F4E78")

    stats = {}
    for ws in wb.worksheets:
        if ws.title == "Index":
            continue
        cats = rows = answered_rows = 0
        per = {label: 0 for label in labels}
        for section in find_sections(ws):
            seen_category = None
            for row in section["rows"]:
                rows += 1
                category = str(ws.cell(row, 2).value or "").strip()
                if category != seen_category:
                    cats += 1
                    seen_category = category
                hit = False
                for label in labels:
                    col = section["out_cols"].get(label)
                    if not col:
                        continue
                    value = ws.cell(row, col).value
                    if value is not None and str(value).strip():
                        per[label] += 1
                        hit = True
                if hit:
                    answered_rows += 1
        stats[ws.title] = (cats, rows, answered_rows, per)

    written = 0
    totals = [0, 0, 0, {label: 0 for label in labels}]
    for row in range(4, index.max_row + 2):
        sheet_name = str(index.cell(row, 4).value or "").strip()
        if not sheet_name or sheet_name not in stats:
            continue
        cats, rows_n, answered_rows, per = stats[sheet_name]
        cells = rows_n * len(labels)
        index.cell(row, 5, cats)
        index.cell(row, 6, rows_n)
        index.cell(row, 7, answered_rows)
        for offset, label in enumerate(labels):
            index.cell(row, 8 + offset, per[label])
        index.cell(row, 12, round(sum(per.values()) / cells * 100, 1) if cells else 0)
        totals[0] += cats
        totals[1] += rows_n
        totals[2] += answered_rows
        for label in labels:
            totals[3][label] += per[label]
        written += 1

    if written:
        row = 4 + written + 1
        index.cell(row, 3, "TOTAL").font = Font(bold=True)
        index.cell(row, 5, totals[0]).font = Font(bold=True)
        index.cell(row, 6, totals[1]).font = Font(bold=True)
        index.cell(row, 7, totals[2]).font = Font(bold=True)
        for offset, label in enumerate(labels):
            index.cell(row, 8 + offset, totals[3][label]).font = Font(bold=True)
        all_cells = totals[1] * len(labels)
        index.cell(row, 12, round(sum(totals[3].values()) / all_cells * 100, 1)
                   if all_cells else 0).font = Font(bold=True)
    for col, width in {7: 14, 8: 10, 9: 10, 10: 13, 11: 8, 12: 9}.items():
        index.column_dimensions[get_column_letter(col)].width = width
    return written


def try_save(wb, path, json_path):
    """
    Save the workbook, tolerating the file being open elsewhere.

    On Windows the temp-file swap fails with PermissionError while Excel has the workbook
    open. That used to end the run. Nothing is lost when it happens: every answer is
    already in the JSONL, and the whole workbook lives in memory here, so the next
    successful save writes all of it. Returns True when it actually landed.
    """
    try:
        refresh_index(wb)  # keep the Index's counts honest on every save
        save_atomically(path, lambda p: wb.save(p))
        write_json(wb, path, json_path)
        return True
    except PermissionError:
        log(f"    {path.name} is open in Excel, so it could not be saved. Close it and "
            "the next answer will write everything; nothing is lost (it is all in "
            f"{json_path.with_name(json_path.stem + '_answers.jsonl').name}).")
        return False
    except Exception as exc:
        log(f"    save failed ({exc.__class__.__name__}: {str(exc)[:90]}); "
            "answers are held in memory and the JSONL, will retry on the next one")
        return False


def build_json(wb, workbook_path):
    """
    Nested exactly as the sheets are organised:
      sub-industry -> tier (the source column heading) -> category -> query -> answers
    """
    data = {
        "generated": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_workbook": workbook_path.name,
        "engines": [e.label for e in llm_engines.ENGINES.values()],
        "totals": {"sub_industries": 0, "categories": 0, "queries": 0, "answers": {}},
        "sub_industries": [],
    }
    answered = {}

    for ws in wb.worksheets:
        if ws.title == "Index":
            continue
        sheet_entry = {"sub_industry": str(ws.cell(1, 1).value or ws.title)
                       .split(" — LLM Response")[0].strip(),
                       "sheet": ws.title, "tiers": []}
        for section in find_sections(ws):
            tier_entry = {"tier": section["tier"], "categories": []}
            current = None
            for row in section["rows"]:
                category = str(ws.cell(row, 2).value or "").strip()
                query = str(ws.cell(row, section["query_col"]).value or "").strip()
                if current is None or current["category"] != category:
                    current = {"category": category, "queries": []}
                    tier_entry["categories"].append(current)
                    data["totals"]["categories"] += 1

                answers = {}
                for label, col in section["out_cols"].items():
                    value = ws.cell(row, col).value
                    text = str(value).strip() if value is not None else ""
                    src_col = section["src_cols"].get(label)
                    src = ws.cell(row, src_col).value if src_col else None
                    src = str(src).strip() if src is not None else ""
                    answers[label] = {"output": text or None, "sources": src or None}
                    if text:
                        answered[label] = answered.get(label, 0) + 1
                current["queries"].append({
                    "n": ws.cell(row, 1).value,
                    "row": row,
                    "query": query,
                    "answers": answers,
                })
                data["totals"]["queries"] += 1
            if tier_entry["categories"]:
                sheet_entry["tiers"].append(tier_entry)
        if sheet_entry["tiers"]:
            data["sub_industries"].append(sheet_entry)
            data["totals"]["sub_industries"] += 1

    data["totals"]["answers"] = answered
    return data


def write_json(wb, workbook_path, json_path):
    payload = build_json(wb, workbook_path)
    save_atomically(json_path, lambda p: p.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"))
    return payload["totals"]


def append_jsonl(path, task, answer, sources=""):
    """
    One line per answer, written the moment it arrives, never rewritten -- and always the
    complete text, even when the copy in the sheet had to be trimmed for Excel.
    """
    record = {
        "ts": datetime.datetime.now().isoformat(timespec="seconds"),
        "sub_industry": task["sheet"], "tier": task["tier"],
        "category": task["category"], "row": task["row"],
        "engine": task["engine"], "query": task["query"],
        "chars": len(answer), "answer": answer,
        "source_count": len([l for l in sources.splitlines()[1:] if l.strip()]),
        "sources": sources or None,
    }
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")


def check_writable(path):
    if not path.exists():
        sys.exit(f"Workbook not found: {path}\nRun build_trackers.py first.")
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        sys.exit(f"{path.name} is open in Excel (or locked). Close it and run again -- "
                 "answers cannot be written while it is open.")


# ----------------------------------------------------------------------------------
# Running
# ----------------------------------------------------------------------------------
def preflight(page, human, engines, args):
    """
    Check each engine once before the run: land on its page, clear any Cloudflare check,
    and confirm a composer is actually there. Front-loading the friction beats finding it
    forty queries in -- and a solved challenge leaves a cf_clearance cookie in the
    profile that serves the rest of the run.

    Returns the engines that are ready. With --solve-challenges it pauses so you can
    click a checkbox yourself.
    """
    print("\nPre-flight check:")
    ready, notes = [], []
    for engine in engines:
        url = (engine.url.format(query="test") if "{query}" in engine.url
               else engine.url)
        status = "?"
        try:
            page.goto(url, wait_until="domcontentloaded")
            llm_engines.settle(page)
            human.pause(0.6, 1.4)

            done = engine.composer or ["#search", "#rso", "form[action='/search']"]
            if not cgq.clear_challenge(page, human, done=done,
                                       manual=args.solve_challenges):
                status = "BLOCKED by a Cloudflare check"
            else:
                landed = llm_engines.safe(lambda: page.url, "?")
                host = landed.split("/")[2] if "//" in landed else "?"
                want = url.split("/")[2]
                if engine.kind == "search":
                    found, _ = llm_engines.first_present(page, done, timeout=8)
                    status = "ready" if found is not None else "no search box"
                elif host != want:
                    status = f"redirected to {host} -- sign in needed"
                else:
                    found, selector = llm_engines.first_present(
                        page, engine.composer, timeout=12)
                    if found is not None:
                        status = "ready"
                        notes.append(f"    {engine.label}: composer {selector}")
                    elif llm_engines.first_present(page, engine.login)[0] is not None:
                        status = "signed out"
                    else:
                        status = "no composer found"
        except Exception as exc:
            status = f"{exc.__class__.__name__}: {str(exc)[:60]}"

        # Whether this profile is warm for the site. A Cloudflare check that keeps
        # coming back usually means cf_clearance was earned in a different browser or
        # profile than the one running here, so the cookie never arrived.
        cookies = llm_engines.safe(lambda: page.context.cookies(url), []) or []
        names = {c.get("name", "") for c in cookies}
        warmth = []
        if "cf_clearance" in names:
            warmth.append("cf_clearance")
        if any(n.startswith(("session", "__Secure", "sessionKey", "auth")) for n in names):
            warmth.append("session")
        cookie_note = ", ".join(warmth) if warmth else f"{len(cookies)} cookies, "
        if not warmth:
            cookie_note += "no cf_clearance/session"

        print(f"  {engine.label:<15} {status:<34} [{cookie_note}]")
        if status == "ready":
            ready.append(engine)

    for note in notes:
        print(note)
    if len(ready) < len(engines):
        missing = [e.label for e in engines if e not in ready]
        print(f"\n  Not ready: {', '.join(missing)}")
        print("  Fix with:  python chatgpt_query.py --login     (sign in to each site)")
        print("             python llm_engines.py --probe <key> (see selector detail)")
        print("             --solve-challenges                  (click checks yourself)")
    return ready


def recover_page(page, human, ctx, args):
    """
    Replace a tab whose session has died. Symptoms are Playwright's "Target page, context
    or browser has been closed" and, on a rebrowser build, "CRSession.__re__getMainWorld
    ... type: 'closed'" -- the main-world handshake failing on a page that navigated or
    closed underneath it. A dead session never recovers on its own, and every later query
    on that page fails the same way, so swap in a fresh one.

    Returns (page, human) -- unchanged if the page is still alive.
    """
    try:
        if not page.is_closed():
            return page, human
    except Exception:
        pass
    if ctx is None:
        return page, human
    try:
        fresh = ctx.new_page()
        log("    the tab had died; opened a fresh one")
        return fresh, cgq.Human(fresh, wpm=args.wpm, jitter=args.jitter)
    except Exception as exc:
        log(f"    could not reopen a tab ({exc.__class__.__name__}); "
            "the browser itself is probably gone")
        return page, human


def looks_rate_limited(answer):
    low = answer.lower()
    return len(answer) < 400 and any(p in low for p in LIMIT_PHRASES)


def fmt_eta(seconds):
    seconds = int(max(0, seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h}h {m:02d}m" if h else f"{m}m {s:02d}s"


def main():
    ap = argparse.ArgumentParser(
        description="Fill the tracker from every LLM, automatically and resumably.")
    ap.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    ap.add_argument("--json", default=None, help="nested JSON (default: <workbook>.json)")
    ap.add_argument("--jsonl", default=None,
                    help="append-only log (default: <workbook>_answers.jsonl)")
    ap.add_argument("--engines", default="all",
                    help="chatgpt, claude, aio, grok -- comma separated, or all")
    ap.add_argument("--sheet", default=None, help="only sheets whose name contains this")
    ap.add_argument("--tier", default=None, help="small, mid or enterprise")
    ap.add_argument("--limit", type=int, default=0, help="stop after N answers")
    ap.add_argument("--dry-run", action="store_true", help="list the work, touch nothing")
    ap.add_argument("--delay-min", type=float, default=6.0)
    ap.add_argument("--delay-max", type=float, default=20.0)
    ap.add_argument("--break-every", type=int, default=25)
    ap.add_argument("--break-for", type=float, default=180.0)
    ap.add_argument("--retries", type=int, default=1)
    ap.add_argument("--stop-on-error", action="store_true")
    ap.add_argument("--index-only", action="store_true",
                    help="recount the Index sheet and exit; no browser, no queries")
    ap.add_argument("--no-preflight", action="store_true",
                    help="skip the per-engine readiness check before the run")
    ap.add_argument("--no-sources", action="store_true",
                    help="skip capturing the cited links (they are captured by default)")
    ap.add_argument("--solve-challenges", action="store_true",
                    help="pause and let you click a Cloudflare check by hand instead of "
                         "skipping the query (only useful if you are watching)")
    ap.add_argument("--skip-engine-on-error", type=int, default=4,
                    help="drop an engine after this many consecutive failures (0 = never)")
    ap.add_argument("--save-every", type=int, default=1,
                    help="save workbook and JSON after every N answers")

    # Browser options, same meanings as in chatgpt_query.py
    ap.add_argument("-p", "--profile", default=cgq.DEFAULT_PROFILE)
    ap.add_argument("--mode", choices=("own", "copy", "real"), default="own")
    ap.add_argument("--profile-dir", default=None)
    ap.add_argument("--refresh-copy", action="store_true")
    ap.add_argument("--chrome-path", default=None)
    ap.add_argument("--cdp-port", type=int, default=cgq.CDP_PORT)
    ap.add_argument("--force-close", action="store_true")
    ap.add_argument("--close-browser", action="store_true")
    ap.add_argument("--no-stealth", action="store_true")
    ap.add_argument("--runtime-fix", default=None,
                    choices=("addBinding", "alwaysIsolated", "off"),
                    help="how rebrowser hides the Runtime.enable leak. alwaysIsolated survives pages that redirect mid-load, such as Cloudflare interstitials; addBinding hides more but can die on them.")
    ap.add_argument("--wpm", type=int, default=72)
    ap.add_argument("--jitter", type=float, default=1.0)
    ap.add_argument("--timeout", type=int, default=240, help="seconds per answer")
    args = ap.parse_args()

    path = pathlib.Path(args.workbook)
    if not path.is_absolute():
        path = HERE / path
    check_writable(path)

    json_path = pathlib.Path(args.json) if args.json else path.with_suffix(".json")
    if not json_path.is_absolute():
        json_path = HERE / json_path
    jsonl_path = (pathlib.Path(args.jsonl) if args.jsonl
                  else path.with_name(path.stem + "_answers.jsonl"))
    if not jsonl_path.is_absolute():
        jsonl_path = HERE / jsonl_path

    engines = llm_engines.resolve(args.engines)
    by_label = {e.label: e for e in engines}

    log(f"Reading {path.name}")
    wb = openpyxl.load_workbook(path)

    if args.index_only:
        rows = refresh_index(wb)
        if try_save(wb, path, json_path):
            log(f"Index recounted for {rows} sheet(s): {path.name}")
        return

    if args.sheet and not pick_sheets(wb, args.sheet):
        print(f"No sheet matches {args.sheet!r}. Sheets:")
        for ws in wb.worksheets:
            print(f"  {ws.title}")
        sys.exit(1)

    tasks, totals = collect_tasks(wb, engines, args.sheet, args.tier)
    if args.limit:
        tasks = tasks[: args.limit]

    print(f"\nEngines: {', '.join(e.label for e in engines)}")
    print(f"{totals['rows']} query rows in scope | {totals['cells']} answer cells | "
          f"{totals['filled']} already answered | {len(tasks)} to run now")
    per_engine = {}
    for task in tasks:
        per_engine[task["engine"]] = per_engine.get(task["engine"], 0) + 1
    for label, count in per_engine.items():
        print(f"  {label:<16} {count:>6}")
    if totals["missing_cols"]:
        print(f"  note: no column for {sorted(totals['missing_cols'])[:3]}")
    print()

    if not tasks:
        log("Nothing pending.")
        write_json(wb, path, json_path)
        log(f"JSON refreshed: {json_path.name}")
        return

    if args.dry_run:
        for task in tasks[:12]:
            print(f"  [{task['engine']:<13}] {task['sheet'][:22]} r{task['row']} | "
                  f"{task['category'][:24]} | {task['query'][:52]}")
        if len(tasks) > 12:
            print(f"  ... and {len(tasks) - 12} more")
        log("Dry run: nothing sent, nothing written.")
        return

    done = failed = unsaved = 0
    consecutive = {e.label: 0 for e in engines}
    dropped = set()
    started = time.time()

    with cgq.sync_playwright() as p:
        page, close = cgq.open_browser(p, args)
        ctx = llm_engines.safe(lambda: page.context, None)  # for reopening a dead tab
        try:
            human = cgq.Human(page, wpm=args.wpm, jitter=args.jitter)
            cgq.fit_window(page)

            if not args.no_preflight:
                ready = preflight(page, human, engines, args)
                ready_labels = {e.label for e in ready}
                if not ready_labels:
                    log("No engine is ready; nothing was sent. See the notes above.")
                    return
                skipped = [t for t in tasks if t["engine"] not in ready_labels]
                if skipped:
                    log(f"Skipping {len(skipped)} task(s) for engines that are not ready.")
                tasks = [t for t in tasks if t["engine"] in ready_labels]
                print()

            for i, task in enumerate(tasks, start=1):
                label = task["engine"]
                if label in dropped:
                    continue
                engine = by_label[label]
                ws = wb[task["sheet"]]
                log(f"[{i}/{len(tasks)}] {label} | {task['sheet'][:20]} r{task['row']} "
                    f"| {task['category'][:22]} | {task['query'][:44]}...")

                answer = None
                for attempt in range(args.retries + 1):
                    try:
                        answer = llm_engines.ask(engine, page, human, task["query"],
                                                 args.timeout,
                                                 manual_challenge=args.solve_challenges)
                        break
                    except Exception as exc:
                        log(f"    attempt {attempt + 1} failed: "
                            f"{exc.__class__.__name__}: {str(exc)[:130]}")
                        # A closed tab or a dead CDP session poisons every later query,
                        # so replace the page rather than failing the rest of the run.
                        page, human = recover_page(page, human, ctx, args)
                        if attempt < args.retries:
                            human.pause(8.0, 20.0)

                if not answer:
                    failed += 1
                    consecutive[label] += 1
                    if args.skip_engine_on_error and \
                            consecutive[label] >= args.skip_engine_on_error:
                        dropped.add(label)
                        log(f"    dropping {label} for this run after "
                            f"{consecutive[label]} failures in a row "
                            f"(try: python llm_engines.py --probe {engine.key})")
                    if args.stop_on_error:
                        log("Stopping at this failure (--stop-on-error).")
                        break
                    continue

                consecutive[label] = 0

                if looks_rate_limited(answer):
                    log(f"    {label} looks rate-limited: {answer[:130]}")
                    dropped.add(label)
                    log(f"    dropping {label} for this run; rerun later to resume")
                    continue

                # Sources come from the page we are still on, so they must be read now,
                # before anything navigates away.
                sources = ""
                if not args.no_sources:
                    sources = llm_engines.extract_sources(page, engine)

                cell = ws.cell(task["row"], task["out_col"],
                               llm_engines.for_cell(answer))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if sources and task["src_col"]:
                    src_cell = ws.cell(task["row"], task["src_col"],
                                       llm_engines.for_cell(sources))
                    src_cell.alignment = Alignment(wrap_text=True, vertical="top")
                done += 1

                # Immediate and crash-proof, and always the untrimmed text.
                append_jsonl(jsonl_path, task, answer, sources)
                if args.save_every <= 1 or done % args.save_every == 0:
                    saved_ok = try_save(wb, path, json_path)
                    if not saved_ok:
                        unsaved += 1
                col_letter = openpyxl.utils.get_column_letter(task["out_col"])
                log(f"    saved {len(answer):,} chars to "
                    f"{task['sheet']}!{col_letter}{task['row']} (+ json, jsonl)")

                if dropped and len(dropped) == len(engines):
                    log("Every engine has been dropped; stopping.")
                    break

                if i < len(tasks):
                    if args.break_every and done and done % args.break_every == 0:
                        log(f"    pausing {args.break_for:.0f}s after {done} answers")
                        time.sleep(args.break_for)
                    else:
                        human.read(answer)
                        time.sleep(random.uniform(args.delay_min, args.delay_max))

                if done >= 3:
                    per = (time.time() - started) / done
                    left = len(tasks) - i
                    log(f"    {per:.0f}s/answer | {left} left | ETA {fmt_eta(per * left)}")
        except KeyboardInterrupt:
            log("Interrupted.")
        finally:
            # Retry the final save a few times: if the workbook is open in Excel, a few
            # seconds is often all it takes for the person to close it.
            for attempt in range(3):
                if try_save(wb, path, json_path):
                    log("Saved.")
                    unsaved = 0
                    break
                if attempt < 2:
                    log("    retrying the save in 10s -- close the workbook if it is open")
                    time.sleep(10)
            else:
                log(f"COULD NOT SAVE {path.name}. Every answer from this run is still in "
                    f"{jsonl_path.name}; close the workbook and rerun to redo only what "
                    "is missing.")
            try:
                close()
            except Exception:
                pass

    print()
    log(f"{done} answered, {failed} failed.")
    if unsaved:
        log(f"{unsaved} save(s) were skipped because the workbook was open; the last "
            "successful save still contains every answer held in memory.")
    if dropped:
        log(f"dropped this run: {', '.join(sorted(dropped))}")
    log(f"Workbook: {path}")
    log(f"JSON:     {json_path}")
    log(f"JSONL:    {jsonl_path}")
    log("Run the same command again to continue where this left off.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nInterrupted.")
